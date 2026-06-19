"""
Importador da ListagemB3.

Lê um ou mais arquivos .xlsx baixados do sistema interno da B3 (cada um
trazendo uma série de uma emissão) e monta uma base única de
investidores para a assembleia. Não soma nada entre arquivos: cada
linha de cada arquivo entra como um registro independente — o usuário
decide manualmente, na conciliação, se duas linhas do mesmo CPF/CNPJ
votam juntas ou separadas (ver conversa de arquitetura: contas
diferentes do mesmo participante podem ter tratamento distinto).

Aba SINTÉTICOS  → vira a base "PARTICIPANTES" (participantes diretos).
Aba APROVADOS   → vira a base "COMITENTES" (investidores finais).
"""
from dataclasses import dataclass, field
from pathlib import Path
from openpyxl import load_workbook

ABA_SINTETICOS = "SINTÉTICOS"
ABA_APROVADOS = "APROVADOS"


@dataclass
class Participante:
    """Uma linha da aba SINTÉTICOS (participante direto — agente custodiante)."""
    razao_social: str
    cnpj: str
    conta: str
    tipo_conta: str
    quantidade: float
    serie: str
    arquivo_origem: str


@dataclass
class Comitente:
    """Uma linha da aba APROVADOS (investidor final)."""
    nome: str
    documento: str          # CPF ou CNPJ, como veio na origem (texto, com zeros)
    tipo_pessoa: str         # "Física" ou "Jurídica"
    quantidade: float
    serie: str
    arquivo_origem: str
    gestor: str | None = None     # preenchido depois, via busca CVM (só PJ)
    status: str | None = None     # "ok" | "fora" | None (pendente)


@dataclass
class BaseImportada:
    participantes: list[Participante] = field(default_factory=list)
    comitentes: list[Comitente] = field(default_factory=list)
    arquivos: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)

    @property
    def total_comitentes(self) -> int:
        return len(self.comitentes)

    @property
    def comitentes_pj(self) -> list[Comitente]:
        return [c for c in self.comitentes if c.tipo_pessoa == "Jurídica"]

    @property
    def comitentes_pf(self) -> list[Comitente]:
        return [c for c in self.comitentes if c.tipo_pessoa == "Física"]


def _texto(valor) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def _numero(valor) -> float:
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    s = str(valor).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _ler_sinteticos(ws, nome_arquivo: str) -> tuple[list[Participante], str | None]:
    participantes = []
    serie = None
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        razao_social = _texto(row[5].value)   # F
        if not razao_social:
            continue
        cod_negociacao = _texto(row[4].value)  # E
        if serie is None:
            serie = cod_negociacao
        participantes.append(Participante(
            razao_social=razao_social,
            cnpj=_texto(row[6].value),         # G
            conta=_texto(row[7].value),        # H
            tipo_conta=_texto(row[8].value),   # I
            quantidade=_numero(row[9].value),  # J
            serie=cod_negociacao,
            arquivo_origem=nome_arquivo,
        ))
    return participantes, serie


def _ler_aprovados(ws, nome_arquivo: str) -> list[Comitente]:
    comitentes = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        nome = _texto(row[8].value)   # I
        if not nome:
            continue
        comitentes.append(Comitente(
            nome=nome,
            documento=_texto(row[9].value),         # J
            tipo_pessoa=_texto(row[10].value),       # K
            quantidade=_numero(row[11].value),       # L
            serie=_texto(row[4].value),              # E
            arquivo_origem=nome_arquivo,
        ))
    return comitentes


def importar_arquivos(caminhos: list[Path]) -> BaseImportada:
    base = BaseImportada()

    for caminho in caminhos:
        caminho = Path(caminho)
        nome_arquivo = caminho.name

        try:
            wb = load_workbook(caminho, data_only=True)
        except Exception as e:
            base.avisos.append(f"Não foi possível abrir {nome_arquivo}: {e}")
            continue

        if ABA_SINTETICOS not in wb.sheetnames:
            base.avisos.append(
                f"{nome_arquivo}: aba '{ABA_SINTETICOS}' não encontrada — pulando."
            )
        else:
            participantes, serie = _ler_sinteticos(wb[ABA_SINTETICOS], nome_arquivo)
            base.participantes.extend(participantes)

        if ABA_APROVADOS not in wb.sheetnames:
            base.avisos.append(
                f"{nome_arquivo}: aba '{ABA_APROVADOS}' não encontrada — pulando."
            )
        else:
            comitentes = _ler_aprovados(wb[ABA_APROVADOS], nome_arquivo)
            base.comitentes.extend(comitentes)

        base.arquivos.append(nome_arquivo)

    return base
