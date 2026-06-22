"""
Gera um Template_Posição_Investidores.xlsx novo a partir da sessão de
assembleia, usando o arquivo modelo real do usuário como base — não um
workbook em branco. Isso preserva 100% da formatação original (cores,
fontes, larguras de coluna, bordas, formato de número) porque o
sistema nunca recria essas células do zero: ele copia o estilo da
linha-modelo (linha 2 de cada aba, que já vem com a formatação certa
e as fórmulas corretas) para cada nova linha de dado.

A aba QUÓRUM não precisa de nenhuma escrita — as fórmulas dos três
blocos (Quantidade Circulação / Quantidade Presentes / Saldo Devedor)
já existem no modelo e continuam funcionando sem alteração.

Cada exportação parte sempre do modelo limpo (nunca de uma exportação
anterior), gerando um arquivo novo — decisão confirmada: cada
assembleia tem população e quantidades diferentes.
"""
import re
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.dashboard.b3_importer import BaseImportada
from app.dashboard.paths import caminho_recurso

MODELO_PATH = caminho_recurso("assets", "Template_Posição_Investidores_modelo.xlsx")

LINHA_MODELO = 2  # linha do template que tem a formatação/fórmulas de referência

_REF_CELULA_RE = re.compile(r"(\$?[A-Z]{1,2})(\d+)\b")


def _copiar_estilo_linha(ws: Worksheet, linha_origem: int, linha_destino: int, max_col: int):
    """Copia fonte, preenchimento, borda, alinhamento e formato de número
    de cada célula da linha_origem para a linha_destino."""
    for col in range(1, max_col + 1):
        origem = ws.cell(row=linha_origem, column=col)
        destino = ws.cell(row=linha_destino, column=col)
        destino.font = copy(origem.font)
        destino.fill = copy(origem.fill)
        destino.border = copy(origem.border)
        destino.alignment = copy(origem.alignment)
        destino.number_format = origem.number_format


def _copiar_formula_ajustada(formula: str, linha_origem: int, linha_destino: int) -> str:
    """Ajusta referências relativas de linha dentro de uma fórmula ao
    copiá-la para outra linha (ex: D2 -> D3 quando a fórmula referencia
    a própria linha). Não toca referências de coluna inteira ($G:$G)
    nem fórmulas que referenciam outra aba com linha fixa (QUÓRUM!$H$4)."""

    def substituir(m):
        col, lin = m.group(1), m.group(2)
        if lin == str(linha_origem):
            return f"{col}{linha_destino}"
        return m.group(0)

    return _REF_CELULA_RE.sub(substituir, formula)


def _mapear_modelo(ws: Worksheet, max_col: int) -> tuple[dict, dict]:
    """Retorna (formulas, valores_padrao) da linha-modelo: fórmulas são
    strings que começam com '=' (precisam de ajuste de referência ao
    copiar); valores_padrao são constantes (ex: a coluna 'Séries' com
    valor fixo 'única') que devem ser replicados em toda linha nova."""
    formulas, valores_padrao = {}, {}
    for col in range(1, max_col + 1):
        valor = ws.cell(row=LINHA_MODELO, column=col).value
        if valor is None:
            continue
        if isinstance(valor, str) and valor.startswith("="):
            formulas[col] = valor
        else:
            valores_padrao[col] = valor
    return formulas, valores_padrao


def _preencher_participantes(ws: Worksheet, base: BaseImportada):
    max_col = ws.max_column
    formulas_modelo, valores_padrao = _mapear_modelo(ws, max_col)
    # colunas que o importador sempre preenche explicitamente — não
    # devem ser sobrescritas pelos valores padrão do modelo
    colunas_dados = {2, 3, 4, 5, 6}

    for i, p in enumerate(base.participantes, start=LINHA_MODELO):
        if i > LINHA_MODELO:
            _copiar_estilo_linha(ws, LINHA_MODELO, i, max_col)
            for col, valor in valores_padrao.items():
                if col not in colunas_dados:
                    ws.cell(row=i, column=col, value=valor)

        ws.cell(row=i, column=2, value=p.razao_social)
        ws.cell(row=i, column=3, value=int(p.cnpj) if p.cnpj.isdigit() else p.cnpj)
        ws.cell(row=i, column=4, value=p.conta)
        ws.cell(row=i, column=5, value=p.tipo_conta)
        ws.cell(row=i, column=6, value=p.quantidade)

        for col, formula in formulas_modelo.items():
            ws.cell(row=i, column=col, value=_copiar_formula_ajustada(formula, LINHA_MODELO, i))


def _preencher_comitentes(ws: Worksheet, base: BaseImportada):
    max_col = ws.max_column
    formulas_modelo, valores_padrao = _mapear_modelo(ws, max_col)
    colunas_dados = {3, 4, 5, 6, 7, 9, 17}

    for i, c in enumerate(base.comitentes, start=LINHA_MODELO):
        if i > LINHA_MODELO:
            _copiar_estilo_linha(ws, LINHA_MODELO, i, max_col)
            for col, valor in valores_padrao.items():
                if col not in colunas_dados:
                    ws.cell(row=i, column=col, value=valor)

        ws.cell(row=i, column=3, value=c.gestor or "")
        ws.cell(row=i, column=4, value=c.nome)
        doc = c.documento
        ws.cell(row=i, column=5, value=int(doc) if doc.isdigit() else doc)
        ws.cell(row=i, column=6, value=c.tipo_pessoa)
        ws.cell(row=i, column=7, value=c.quantidade)
        if c.status:
            ws.cell(row=i, column=9, value=c.status.upper())
        ws.cell(row=i, column=17, value=c.serie)

        for col, formula in formulas_modelo.items():
            ws.cell(row=i, column=col, value=_copiar_formula_ajustada(formula, LINHA_MODELO, i))


def gerar_template(base: BaseImportada, pasta_destino: Path) -> Path:
    """Gera um novo Template_Posição_Investidores.xlsx a partir do
    modelo real do usuário, preenchendo PARTICIPANTES e COMITENTES.
    A aba QUÓRUM não é tocada — suas fórmulas já existem no modelo."""
    if not MODELO_PATH.exists():
        raise FileNotFoundError(
            f"Arquivo modelo não encontrado em {MODELO_PATH}. "
            "Verifique se app/dashboard/assets/Template_Posição_Investidores_modelo.xlsx existe."
        )

    wb = load_workbook(MODELO_PATH, data_only=False)

    _preencher_participantes(wb["PARTICIPANTES"], base)
    _preencher_comitentes(wb["COMITENTES"], base)

    pasta_destino = Path(pasta_destino)
    pasta_destino.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho_saida = pasta_destino / f"Template_Posição_Investidores_{timestamp}.xlsx"

    wb.save(caminho_saida)
    return caminho_saida
