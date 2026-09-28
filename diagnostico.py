"""
Diagnóstico: por que o número de votos não bate com o número de pastas?

Uso (na raiz do projeto, com o venv ativo):
    python diagnostico.py "C:\\caminho\\pasta_raiz" "NOME_DA_PASTA_IGNORADA" [planilha.xlsx]

Funciona tanto nas pastas ORIGINAIS (PDFs sem renomear) quanto numa pasta com
os PDFs já renomeados. Quando o nome não está no padrão "CPF - votos", o script
LÊ o PDF (sem renomear nada) com o modelo "Instrução [X]" para descobrir CPF/votos.

Gera diagnostico.csv na pasta raiz com uma linha por PDF.
"""
import csv
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

def ler_conteudo(p: Path):
    """Extrai CPF/CNPJ e votos do conteúdo, sem alterar o arquivo."""
    try:
        import pdfplumber
        from app.models.instrucao.instrucao_colchete import InstrucaoColchete
        with pdfplumber.open(p) as pdf:
            m = InstrucaoColchete(pdf)
            return m.extrair_cpf() or "", ", ".join(m.extrair_votos()) or "SEM_VOTO"
    except Exception as e:
        return "", f"ERRO: {e}"


NOME_RE = re.compile(r"^(\d{11}|\d{14})\s*-\s*(.+)$")


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return
    raiz = Path(sys.argv[1])
    ignorada = sys.argv[2].lower()
    excel = Path(sys.argv[3]) if len(sys.argv) > 3 else None

    pdfs = sorted(raiz.rglob("*.pdf"))
    linhas = []
    for p in pdfs:
        rel = p.relative_to(raiz)
        m = NOME_RE.match(p.stem.strip())
        if m:
            doc, votos, origem = m.group(1), m.group(2), "nome"
        else:
            doc, votos = ler_conteudo(p)
            origem = "conteudo"
        linhas.append({
            "pasta_nivel1": rel.parts[0] if len(rel.parts) > 1 else "(raiz)",
            "profundidade": len(rel.parts) - 1,
            "ignorada": any(ignorada in part.lower() for part in rel.parts[:-1]),
            "doc": doc,
            "tipo_doc": ("CNPJ" if len(doc) == 14 else "CPF") if doc else "SEM_DOC",
            "votos": votos,
            "origem": origem,
            "modificado": datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M"),
            "arquivo": str(rel),
        })

    validos = [l for l in linhas if not l["ignorada"]]
    pastas = {l["pasta_nivel1"] for l in validos}
    por_pasta = Counter(l["pasta_nivel1"] for l in validos)
    por_doc = defaultdict(list)
    for l in validos:
        if l["doc"]:
            por_doc[l["doc"]].append(l)

    n_ign = sum(1 for l in linhas if l["ignorada"])
    if n_ign:
        print(f"\n⚠ Se você copiou tudo com a busca *.pdf do Explorer a partir da raiz,"
              f" estes {n_ign} PDF(s) da pasta ignorada foram copiados junto:")
        for l in linhas:
            if l["ignorada"]:
                print(f"   {l['doc'] or '?'}  [{l['votos']}]  {l['arquivo']}")
    print(f"\nPDFs totais (rglob, inclui subpastas): {len(linhas)}")
    print(f"  dentro da pasta ignorada:            {len(linhas) - len(validos)}")
    print(f"  fora dela:                           {len(validos)}")
    print(f"Pastas de 1º nível com PDF (fora da ignorada): {len(pastas)}")
    print(f"CPFs/CNPJs distintos (fora da ignorada):       {len(por_doc)}")

    multi = {k: v for k, v in por_pasta.items() if v > 1}
    if multi:
        print(f"\n▶ Pastas com mais de 1 PDF ({len(multi)}):")
        for pasta, n in sorted(multi.items()):
            print(f"   {n}x  {pasta}")

    dups = {k: v for k, v in por_doc.items() if len(v) > 1}
    if dups:
        print(f"\n▶ Mesmo CPF/CNPJ em vários PDFs ({len(dups)}) — o Excel usa só UM deles:")
        for doc, ls in dups.items():
            print(f"   {doc}")
            for l in sorted(ls, key=lambda x: x["modificado"]):
                print(f"      {l['modificado']}  [{l['votos']}]  {l['arquivo']}")

    fora = [l for l in validos if l["tipo_doc"] == "SEM_DOC"]
    if fora:
        print(f"\n▶ PDFs sem CPF/CNPJ identificado ({len(fora)}):")
        for l in fora:
            print(f"   {l['arquivo']}  [{l['votos']}]")

    cnpjs = [l for l in validos if l["tipo_doc"] == "CNPJ"]
    if cnpjs:
        print(f"\n▶ Arquivos com CNPJ (a versão atual do Excel NÃO lê 14 dígitos) ({len(cnpjs)}):")
        for l in cnpjs:
            print(f"   {l['arquivo']}")

    fundas = [l for l in validos if l["profundidade"] > 1]
    if fundas:
        print(f"\n▶ PDFs em subpastas de 2º nível ou mais ({len(fundas)}):")
        for l in fundas:
            print(f"   {l['arquivo']}")

    if excel:
        from openpyxl import load_workbook
        ws = load_workbook(excel, read_only=True)["COMITENTES"]
        linhas_cpf = Counter()
        for (cpf,) in ws.iter_rows(min_row=2, min_col=5, max_col=5, values_only=True):
            if cpf is None:
                break
            linhas_cpf[re.sub(r"\D", "", str(cpf)).zfill(11)] += 1
        rep = {d: n for d, n in linhas_cpf.items() if n > 1 and d in por_doc}
        print(f"\nPlanilha: {sum(linhas_cpf.values())} linhas, {len(linhas_cpf)} CPFs/CNPJs distintos")
        if rep:
            extra = sum(n - 1 for n in rep.values())
            print(f"▶ {len(rep)} votante(s) aparecem em várias linhas (várias contas/custodiantes)"
                  f" → +{extra} linha(s) preenchida(s) além do nº de pessoas:")
            for d, n in rep.items():
                print(f"   {d}: {n} linhas")
        sem_linha = [d for d in por_doc if d not in linhas_cpf]
        if sem_linha:
            print(f"▶ {len(sem_linha)} CPF/CNPJ dos PDFs que não estão na planilha: {', '.join(sem_linha)}")

    saida = raiz / "diagnostico.csv"
    with open(saida, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(linhas[0].keys()) if linhas else ["arquivo"], delimiter=";")
        w.writeheader()
        w.writerows(linhas)
    print(f"\nDetalhe completo: {saida}")


if __name__ == "__main__":
    main()
